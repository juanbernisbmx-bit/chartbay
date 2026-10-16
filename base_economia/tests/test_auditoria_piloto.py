"""Contrasta la instantanea con celdas concretas de los originales conservados."""
import hashlib
import json
import sys
import unittest
from pathlib import Path
from unittest import mock

import openpyxl
import xlrd

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts import preparar_piloto, top_provincial

ORIGINALS = ROOT / "datos" / "originales"
IPC = {
    "ipc_ng_nacional": ("sh_ipc_09_26.xls", "Índices IPC Cobertura Nacional", "Total nacional", 6, 10, "Argentina"),
    "ipc_ng_gba": ("sh_ipc_aperturas.xls", "Índices aperturas", "Región GBA", 6, 8, "Gran Buenos Aires"),
    "ipc_ng_pampeana": ("sh_ipc_aperturas.xls", "Índices aperturas", "Región Pampeana", 56, 58, "Pampeana"),
    "ipc_ng_noa": ("sh_ipc_aperturas.xls", "Índices aperturas", "Región Noroeste", 104, 106, "Noroeste"),
    "ipc_ng_nea": ("sh_ipc_aperturas.xls", "Índices aperturas", "Región Noreste", 152, 154, "Noreste"),
    "ipc_ng_cuyo": ("sh_ipc_aperturas.xls", "Índices aperturas", "Región Cuyo", 200, 202, "Cuyo"),
    "ipc_ng_patagonia": ("sh_ipc_aperturas.xls", "Índices aperturas", "Región Patagonia", 248, 250, "Patagonia"),
}
TOP = (
    (2024, "top_mensual_2024_0.xlsx", 10, 12, "2025-08-19"),
    (2025, "top_mensual_2025_12.xlsx", 11, 12, "2026-09-14"),
    (2026, "top_mensual_2026_18.xlsx", 11, 8, "2026-10-02"),
)


class AuditTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.bundle = json.loads((ROOT / "datos" / "piloto" / "catalogo.json").read_text(encoding="utf-8"))
        cls.series = {item["id"]: item for item in cls.bundle["series"] if item["available"]}

    def test_every_ipc_observation_matches_its_source_cell(self):
        checked = 0
        for series_id, (filename, sheet_name, heading, heading_row, value_row, territory) in IPC.items():
            with self.subTest(series=series_id):
                series = self.series[series_id]
                original = ORIGINALS / filename
                book = xlrd.open_workbook(original)
                sheet = book.sheet_by_name(sheet_name)
                self.assertEqual(sheet.cell_value(heading_row - 1, 0), heading)
                self.assertEqual(sheet.cell_value(value_row - 1, 0), "Nivel general")
                self.assertEqual(series["source_row"], value_row)
                self.assertEqual(series["territory"], territory)
                self.assertEqual(series["frequency"], "Mensual")
                self.assertEqual(series["unit"], "Índice")
                self.assertIn("Diciembre 2016=100", series["basis"])
                self.assertEqual(series["source_id"], "INDEC_IPC")
                self.assertEqual(series["original"]["sha256"], hashlib.sha256(original.read_bytes()).hexdigest())
                expected = []
                for column in range(1, sheet.ncols):
                    header = sheet.cell(heading_row - 1, column)
                    if header.ctype == xlrd.XL_CELL_EMPTY:
                        continue
                    self.assertEqual(header.ctype, xlrd.XL_CELL_DATE)
                    period = xlrd.xldate_as_datetime(header.value, book.datemode)
                    value = sheet.cell(value_row - 1, column)
                    self.assertEqual(value.ctype, xlrd.XL_CELL_NUMBER)
                    expected.append({"date": period.strftime("%Y-%m-01"), "value": value.value})
                self.assertEqual(series["observations"], expected)
                self.assertEqual(len(expected), 117)
                checked += len(expected)
        self.assertEqual(checked, 819)

    def test_every_top_observation_matches_cordoba_source_cell(self):
        series = self.series["top_mensual_cordoba"]
        self.assertEqual(series["source_row"], None)
        self.assertEqual(series["territory"], "Córdoba")
        self.assertEqual(series["frequency"], "Mensual")
        self.assertEqual(series["unit"], "Millones de ARS corrientes")
        self.assertEqual(series["source_id"], "TOP_MENSUAL")
        self.assertEqual([item["file"] for item in series["originals"]], [item[1] for item in TOP])
        expected = []
        for year, filename, row_number, count, updated_on in TOP:
            with self.subTest(year=year):
                original = ORIGINALS / filename
                provenance = next(item for item in series["originals"] if item["file"] == filename)
                self.assertEqual(provenance["sha256"], hashlib.sha256(original.read_bytes()).hexdigest())
                self.assertEqual(provenance["source_row"], row_number)
                self.assertIsNone(provenance["published_at"])
                book = openpyxl.load_workbook(original, read_only=True, data_only=True)
                try:
                    sheet = book["TOTAL"]
                    self.assertEqual(sheet.cell(row=3, column=2).value, year)
                    self.assertEqual(sheet.cell(row=4, column=2).value, "en millones de pesos")
                    self.assertEqual(sheet.cell(row=1, column=15).value.strftime("%Y-%m-%d"), updated_on)
                    self.assertEqual(sheet.cell(row=row_number, column=2).value, "Córdoba")
                    values = [sheet.cell(row=row_number, column=month + 2).value for month in range(1, 13)]
                    self.assertTrue(all(isinstance(value, (int, float)) and value > 0 for value in values[:count]))
                    self.assertTrue(all(value in (0, None) for value in values[count:]))
                    total = sheet.cell(row=row_number, column=15).value
                    self.assertAlmostEqual(sum(values[:count]), total, delta=1e-6)
                    expected.extend({"date": f"{year}-{month:02d}-01", "value": value}
                                    for month, value in enumerate(values[:count], 1))
                finally:
                    book.close()
        self.assertEqual(series["observations"], expected)
        self.assertEqual(len(expected), 32)
        self.assertEqual(sum(item["count"] for item in self.series.values()), 851)

    def test_snapshot_id_changes_when_extracted_value_changes(self):
        baseline = preparar_piloto.build()["release_id"]
        originals, series = top_provincial.load()
        series["observations"][0]["value"] += 1
        with mock.patch.object(top_provincial, "load", return_value=(originals, series)):
            self.assertNotEqual(preparar_piloto.build()["release_id"], baseline)
        self.assertEqual(self.bundle["release_id"], baseline)


if __name__ == "__main__":
    unittest.main()
